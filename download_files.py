"""
download_files.py
Download every direct-link file on the workbook's Files tab (or a CSV saved from it) and save each one
as <files_dir>/<file_name>, i.e. in the right folder with the right name.

Handles:
  * links that go straight to a PDF (or that start a PDF download instead of opening a page)
  * Google Drive share links (…/file/d/<id>/view), including Drive's "can't scan for viruses" step
  * article / landing pages: finds the PDF through the page's citation_pdf_url tag, or through a single
    PDF link on the page (those are marked "ok-check" so you can glance at them)
  * if the "Download from" link fails, tries the "DOI link" column next
  * if plain downloading is blocked, retries in a real (hidden) browser

Skips the Infogram / Looker Studio / Canva rows: those are web pages, captured by capture_web_reports.py.
Anything it can't get is listed as "manual" in the log with the reason.

SETUP (once):
    pip install requests playwright openpyxl pillow pypdf
    python -m playwright install chromium

USE
    python download_files.py --workbook "2026_10-22 IVDRS_SUDORS_Reconciled_Deposit_List.xlsx" --files-dir files
    python download_files.py --csv my_rows.csv --files-dir files
    python download_files.py --workbook "...xlsx" --sheet "Not Open Access Files" --files-dir files
Options:
    --only IVSU-0013 IVSU-0045   just these records
    --overwrite                  re-download files that already exist (default: skip them)
    --no-browser                 don't use the browser fallback (faster, gets fewer files)
    --log path.csv               where to write the results log (default: download_log.csv)
"""
import argparse
import csv
import re
import sys
import time
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urljoin, urlparse, parse_qs, unquote

import requests

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0 Safari/537.36")
HEADERS = {"User-Agent": UA, "Accept": "application/pdf,text/html;q=0.9,*/*;q=0.8"}
SKIP_TYPES = {"Infogram web visualization", "Looker Studio dashboard", "Canva design (view link)"}
PAUSE = 1.0          # seconds between records, to be polite to servers
TIMEOUT = 60


# ---------------------------------------------------------------- input rows
def read_rows(path, sheet="Files"):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            return [{k.strip(): (v.strip() if isinstance(v, str) else v) for k, v in r.items() if k}
                    for r in csv.DictReader(f)]
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    head = [c.value for c in ws[1]]
    return [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True)]


def jobs_from(rows, only=None):
    for r in rows:
        rid, name = r.get("record_id"), r.get("file_name")
        if not rid or not str(rid).startswith("IVSU") or not name:
            continue
        if only and rid not in only:
            continue
        if r.get("Link type") in SKIP_TYPES:
            continue
        urls = [u for u in (r.get("Download from"), r.get("DOI link")) if u]
        yield rid, name, list(dict.fromkeys(urls))


# ---------------------------------------------------------------- helpers
def is_pdf(b):
    return bool(b) and b[:5] == b"%PDF-" and len(b) > 1000


class PageLinks(HTMLParser):
    """Collects citation_pdf_url meta, PDF hrefs, and Google Drive download forms from HTML."""
    def __init__(self):
        super().__init__()
        self.citation, self.pdf_links, self.forms, self._form = None, [], [], None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "meta" and (a.get("name") or "").lower() == "citation_pdf_url" and a.get("content"):
            self.citation = a["content"]
        elif tag == "a" and a.get("href"):
            self.pdf_links.append(a["href"])
        elif tag == "form":
            self._form = {"action": a.get("action", ""), "id": a.get("id", ""), "fields": {}}
        elif tag == "input" and self._form is not None and a.get("name"):
            self._form["fields"][a["name"]] = a.get("value", "")

    def handle_endtag(self, tag):
        if tag == "form" and self._form is not None:
            self.forms.append(self._form); self._form = None


def unwrap(href):
    """Turn viewer/redirect wrappers (docs.google.com/viewer?url=…, google.com/url?q=…) into the real URL."""
    q = parse_qs(urlparse(href).query)
    for key in ("url", "q", "file"):
        if key in q and ".pdf" in q[key][0].lower():
            return unquote(q[key][0])
    return href


def pdf_links_in(html, base):
    p = PageLinks()
    try:
        p.feed(html)
    except Exception:
        pass
    links = []
    for h in p.pdf_links:
        h = unwrap(urljoin(base, h))
        if re.search(r"\.pdf($|[?#])", h, re.I) or "/pdf" in urlparse(h).path.lower():
            links.append(h.split("#")[0])
    cit = urljoin(base, p.citation) if p.citation else None
    return cit, list(dict.fromkeys(links)), p.forms


def drive_id(url):
    m = re.search(r"drive\.google\.com/(?:file/d/|open\?id=|uc\?(?:.*&)?id=)([\w-]{10,})", url)
    return m.group(1) if m else None


# ---------------------------------------------------------------- download strategies
def get_plain(sess, url):
    r = sess.get(url, headers=HEADERS, timeout=TIMEOUT, allow_redirects=True)
    return r.status_code, r.headers.get("content-type", ""), r.content, r.url


def try_drive(sess, url):
    fid = drive_id(url)
    if not fid:
        return None, None
    dl = f"https://drive.usercontent.google.com/download?id={fid}&export=download"
    code, ctype, body, final = get_plain(sess, dl)
    if is_pdf(body):
        return body, "google-drive"
    # Large files: Drive shows a "can't scan for viruses" form; submit it.
    _, _, forms = pdf_links_in(body.decode("utf-8", "ignore"), final)
    for f in forms:
        if f["fields"].get("id") == fid or "download" in (f["id"] + f["action"]):
            r = sess.get(urljoin(final, f["action"] or dl), params=f["fields"], headers=HEADERS, timeout=TIMEOUT)
            if is_pdf(r.content):
                return r.content, "google-drive (confirm step)"
    return None, f"Google Drive returned no PDF (HTTP {code}); file may not be shared publicly"


def try_requests(sess, url):
    """Returns (pdf_bytes, method, note)."""
    if drive_id(url):
        body, msg = try_drive(sess, url)
        return (body, msg, "") if body else (None, None, msg)
    try:
        code, ctype, body, final = get_plain(sess, url)
    except requests.RequestException as e:
        return None, None, f"request error: {e.__class__.__name__}"
    if is_pdf(body):
        return body, "direct", ""
    if code >= 400:
        return None, None, f"HTTP {code}"
    if "html" not in ctype.lower():
        return None, None, f"not a PDF ({ctype or 'unknown type'})"
    cit, links, _ = pdf_links_in(body.decode("utf-8", "ignore"), final)
    if cit:
        try:
            c2, _, b2, _ = get_plain(sess, cit)
            if is_pdf(b2):
                return b2, "page's citation_pdf_url", ""
        except requests.RequestException:
            pass
    if len(links) == 1:
        try:
            c2, _, b2, _ = get_plain(sess, links[0])
            if is_pdf(b2):
                return b2, "only PDF link on the page (check it)", ""
        except requests.RequestException:
            pass
    if len(links) > 1:
        return None, None, f"page has {len(links)} PDF links; pick the right one by hand"
    return None, None, "page has no PDF link"


class Browser:
    """Hidden Chromium for sites that block plain downloads or build the page with JavaScript."""
    def __init__(self):
        from playwright.sync_api import sync_playwright
        self._pw = sync_playwright().start()
        self._b = self._pw.chromium.launch()
        self.ctx = self._b.new_context(accept_downloads=True, user_agent=UA)

    def fetch(self, url):
        page = self.ctx.new_page()
        got = {}
        page.on("download", lambda d: got.setdefault("d", d))
        try:
            try:
                resp = page.goto(url, wait_until="load", timeout=TIMEOUT * 1000)
            except Exception:
                resp = None                      # navigating to a file that downloads raises; that's fine
            page.wait_for_timeout(3000)          # (not time.sleep: Playwright must keep handling events)
            if "d" in got:
                p = got["d"].path()
                body = Path(p).read_bytes() if p else b""
                if is_pdf(body):
                    return body, "browser (download)", ""
            if resp is not None:
                try:
                    body = resp.body()
                    if is_pdf(body):
                        return body, "browser (direct)", ""
                except Exception:
                    pass
            html = page.content()
            cit, links, _ = pdf_links_in(html, page.url)
            for cand, how in ([(cit, "browser + citation_pdf_url")] if cit else []) + \
                             ([(links[0], "browser + only PDF link on the page (check it)")] if len(links) == 1 else []):
                r = self.ctx.request.get(cand, timeout=TIMEOUT * 1000)   # uses the page's cookies
                if is_pdf(r.body()):
                    return r.body(), how, ""
            if len(links) > 1:
                return None, None, f"page has {len(links)} PDF links; pick the right one by hand"
            return None, None, "browser found no PDF (may need sign-in or library access)"
        finally:
            page.close()

    def close(self):
        self._b.close(); self._pw.stop()


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook"); ap.add_argument("--csv"); ap.add_argument("--sheet", default="Files")
    ap.add_argument("--files-dir", default="files")
    ap.add_argument("--only", nargs="*")
    ap.add_argument("--overwrite", action="store_true")
    ap.add_argument("--no-browser", action="store_true")
    ap.add_argument("--log", default="download_log.csv")
    a = ap.parse_args()
    if not (a.workbook or a.csv):
        ap.error("give --workbook or --csv")

    jobs = list(jobs_from(read_rows(a.csv or a.workbook, a.sheet), set(a.only or [])))
    base = Path(a.files_dir)
    print(f"{len(jobs)} file(s) to get (Infogram/Looker/Canva rows skipped: use capture_web_reports.py)")
    sess = requests.Session()
    browser = None
    log = []
    counts = {}
    for n, (rid, name, urls) in enumerate(jobs, 1):
        out = base / name
        if out.exists() and not a.overwrite:
            status, method, used, note, size = "skipped (exists)", "", "", "", out.stat().st_size
        else:
            body = method = used = None; notes = []
            for u in urls:
                body, method, note = try_requests(sess, u)
                if body: used = u; break
                notes.append(note)
            if not body and not a.no_browser:
                if browser is None:
                    try:
                        browser = Browser()
                    except Exception as e:
                        print(f"   (browser unavailable: {e}); continuing without it"); a.no_browser = True
                if browser:
                    for u in urls:
                        if drive_id(u):
                            continue
                        body, method, note = browser.fetch(u)
                        if body: used = u; break
                        notes.append(note)
            if body:
                out.parent.mkdir(parents=True, exist_ok=True)
                tmp = out.with_suffix(".part"); tmp.write_bytes(body); tmp.replace(out)
                status = "ok-check" if "check" in method else "ok"
                note, size = "", len(body)
            else:
                status, used, size = "manual", (urls[0] if urls else ""), 0
                note = "; ".join(dict.fromkeys(x for x in notes if x)) or "no link"
        counts[status] = counts.get(status, 0) + 1
        print(f"[{n:>3}/{len(jobs)}] {status:<16} {rid}  {name}" + (f"   <- {note}" if note else ""))
        log.append({"record_id": rid, "file_name": name, "status": status, "method": method or "",
                    "url_used": used or "", "bytes": size, "note": note})
        time.sleep(PAUSE)
    if browser:
        browser.close()

    with open(a.log, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(log[0].keys()) if log else ["record_id"])
        w.writeheader(); w.writerows(log)
    print("\nSummary: " + ", ".join(f"{k}: {v}" for k, v in sorted(counts.items())))
    print(f"Log written to {Path(a.log).resolve()}")
    print("ok-check = taken from the only PDF link on a page; open it to confirm it's the right document.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
