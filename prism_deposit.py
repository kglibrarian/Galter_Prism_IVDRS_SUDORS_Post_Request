"""Deposit the IVDRS & SUDORS reconciled workbook to Prism (InvenioRDM) — staging first.

Used by 05_Prism_Staging_Deposit.ipynb, 06_Prism_Production_Deposit.ipynb and 07_Prism_Community_Report.ipynb. Reads the workbook tabs directly:
  Metadata for Deposit (Action = Deposit), Creators, Subjects, Files, Community Members.

Safety:
  * The token is read from config.py (or .env) and is never printed, logged or written to disk.
  * Nothing is sent unless the caller passes send=True; dry runs only build payloads.
  * Drafts are created/updated and files uploaded separately from publishing.
  * A registry (record_id -> Prism draft id, files, status) makes every step safe to re-run:
    an existing draft is updated, never duplicated; files already uploaded with the same checksum are skipped.
  * Funding is never sent. Scopus author IDs are never sent (Prism holds ORCIDs only).
"""
import ast, csv, hashlib, html, json, os, re, tempfile, time
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import openpyxl

# ---------------------------------------------------------------- config
def load_config(path):
    """Read NAME = "value" lines from config.py (parsed, not executed) or NAME=value lines from .env."""
    path = Path(path); out = {}
    if not path.exists():
        return out
    text = path.read_text(encoding="utf-8")
    if path.suffix == ".py":
        for node in ast.parse(text).body:
            if isinstance(node, ast.Assign) and isinstance(node.value, ast.Constant) and isinstance(node.value.value, str):
                for t in node.targets:
                    if isinstance(t, ast.Name):
                        out[t.id] = node.value.value
    else:
        for line in text.splitlines():
            m = re.match(r"\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.*)\s*$", line)
            if m and not line.lstrip().startswith("#"):
                out[m.group(1)] = m.group(2).strip().strip('"').strip("'")
    return out


# ---------------------------------------------------------------- workbook
def _rows(ws, header_row=1, first_data_row=2):
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h).strip() if h is not None else None for h in rows[header_row - 1]]
    out = []
    for r in rows[first_data_row - 1:]:
        d = {h: v for h, v in zip(head, r) if h}
        if str(d.get("record_id") or "").startswith("IVSU-") or (d.get("name") and "@" in str(d.get("email") or "")):
            out.append(d)
    return out

def _s(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""

def _split(v, sep="; "):
    return [x.strip() for x in str(v).split(sep)] if v not in (None, "") else []

def read_workbook(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    main = [r for r in _rows(wb["Metadata for Deposit"], 1, 3) if r.get("Action") == "Deposit"]
    creators, subjects, files = {}, {}, {}
    for r in _rows(wb["Creators"]):
        creators.setdefault(r["record_id"], []).append(r)
    for r in _rows(wb["Subjects"]):
        subjects.setdefault(r["record_id"], []).append(r)
    for r in _rows(wb["Files"]):
        files.setdefault(r["record_id"], []).append(r)
    members = _rows(wb["Community Members"]) if "Community Members" in wb.sheetnames else []
    wb.close()      # read-only workbooks keep the file open (and locked on Windows) until closed
    for v in creators.values(): v.sort(key=lambda x: int(x["order"]))
    for v in files.values(): v.sort(key=lambda x: int(x["file_order"]))
    recs = {}
    for r in main:
        rid = r["record_id"]
        recs[rid] = {"row": r, "creators": creators.get(rid, []), "subjects": subjects.get(rid, []),
                     "files": files.get(rid, [])}
    return recs, members


# ---------------------------------------------------------------- payloads
def _person(family, given, orcid=None):
    p = {"type": "personal", "family_name": family, "given_name": given}
    if orcid:
        p["identifiers"] = [{"scheme": "orcid", "identifier": orcid}]
    return p

def _agents(row, prefix, given_col):
    """Creators or contributors from the main sheet (aligned '; ' lists)."""
    kind = _s(row.get(f"{prefix}: Person or Organization")).lower()
    out = []
    if kind.startswith("org"):
        names = _split(row.get(f"{prefix}: Name"))
        affs = _split(row.get(f"{prefix}: Affiliations"))
        roles = _split(row.get(f"{prefix}: Role"))
        for i, n in enumerate(names):
            a = {"person_or_org": {"type": "organizational", "name": n}}
            if affs and len(names) == 1:
                a["affiliations"] = [{"name": x} for x in affs]
            if roles:
                a["_role"] = roles[min(i, len(roles) - 1)]
            out.append(a)
    elif kind.startswith("pers"):
        fam = _split(row.get(f"{prefix}: Family name")); giv = _split(row.get(given_col))
        orc = _split(row.get(f"{prefix}: Name identifiers"))
        affs = _split(row.get(f"{prefix}: Affiliations"))
        roles = _split(row.get(f"{prefix}: Role"))
        for i, f in enumerate(fam):
            a = {"person_or_org": _person(f, giv[i] if i < len(giv) else "", orc[i] if i < len(orc) and orc[i] else None)}
            if len(affs) == len(fam) and affs[i]:
                a["affiliations"] = [{"name": x} for x in affs[i].split(" | ")]
            elif len(affs) == 1 and len(fam) == 1:
                a["affiliations"] = [{"name": affs[0]}]
            if roles:
                a["_role"] = roles[min(i, len(roles) - 1)]
            out.append(a)
    return out

def _description(text):
    paras = [p for p in re.split(r"\n\s*\n", str(text or "").strip()) if p.strip()]
    return "".join(f"<p>{html.escape(_s(p))}</p>" for p in paras)

def build_payload(rec, vocab, include_journal_fields=False):
    """Return (payload, problems). vocab comes from resolve_vocab()."""
    r = rec["row"]; problems = []
    rtype = _s(r.get("Resource type")).lower()
    md = {}
    rt_id = vocab["resource_types"].get(rtype)
    if rt_id: md["resource_type"] = {"id": rt_id}
    else: problems.append(f"resource type '{rtype}' has no Prism id")
    md["title"] = _s(r.get("Title"))
    md["publication_date"] = _s(r.get("Publication Date"))
    if not md["publication_date"]: problems.append("no publication date")

    # creators: articles from the Creators tab (one row per author); everything else from the main sheet
    if rec["creators"]:
        creators = []
        for c in rec["creators"]:
            a = {"person_or_org": _person(_s(c["family_name"]), _s(c["given_names"]), _s(c.get("orcid")) or None)}
            if c.get("affiliations"):
                a["affiliations"] = [{"name": x.strip()} for x in str(c["affiliations"]).split(" | ") if x.strip()]
            creators.append(a)       # scopus_author_id / source are staff-only and never sent
    else:
        creators = _agents(r, "Creators", "Creators: Given name(s)")
    for a in creators: a.pop("_role", None)
    md["creators"] = creators
    if not creators: problems.append("no creators")

    contributors = []
    for a in _agents(r, "Contributors", "Contributors: Given names(s)"):
        role = a.pop("_role", None)
        rid = vocab["roles"].get(role) if role else None
        if role and not rid:
            problems.append(f"contributor role '{role}' not found in Prism; sent without a role")
        if rid: a["role"] = {"id": rid}
        contributors.append(a)
    if contributors: md["contributors"] = contributors

    if r.get("Description"): md["description"] = _description(r["Description"])
    lic = _s(r.get("Licenses: Title"))
    if lic:
        if vocab["licenses"].get(lic):
            md["rights"] = [{"id": vocab["licenses"][lic]}]
        else:
            md["rights"] = [{"title": {"en": _s(r.get("Licenses: Description")) or lic},
                             **({"link": _s(r["Licenses: Link"])} if r.get("Licenses: Link") else {})}]
            problems.append(f"license '{lic}' not in Prism's list; sent as a custom license")
    subj = []
    for s in rec["subjects"]:
        if _s(s.get("include (y/n)")).lower() != "y": continue
        entry = json.loads(s["Prism metadata.subjects entry"])
        if entry not in subj: subj.append(entry)
    if subj: md["subjects"] = subj
    langs = _split(r.get("Languages"))
    if langs: md["languages"] = [{"id": x} for x in langs]
    if r.get("Publisher"): md["publisher"] = _s(r["Publisher"])
    ids = []
    for ident, scheme in zip(_split(r.get("Alternate Identifiers: Identifier")), _split(r.get("Alternate identifiers: Scheme"))):
        ids.append({"identifier": ident, "scheme": scheme.lower()})
    if ids: md["identifiers"] = ids

    payload = {"metadata": md, "access": {"record": "public", "files": "public"},
               "files": {"enabled": bool(rec["files"])}}
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", _s(r.get("DOI")), flags=re.I)
    if rtype == "journal article":
        if doi:
            payload["pids"] = {"doi": {"identifier": doi, "provider": "external"}}
        else:
            problems.append("journal article without a publisher DOI")
    else:
        payload["pids"] = {}
    if include_journal_fields and rtype == "journal article" and r.get("Journal (vol./no./pp.)"):
        parts = _split(r["Journal (vol./no./pp.)"]); j = {"title": parts[0]}
        for p in parts[1:]:
            for k, pre in (("volume", "vol. "), ("issue", "no. "), ("pages", "pp. ")):
                if p.startswith(pre): j[k] = p[len(pre):]
        payload["custom_fields"] = {"journal:journal": j}
    return payload, problems


def check_files(recs, files_dir):
    """Return {record_id: [(key, path, problem_or_None)]}."""
    out = {}
    for rid, rec in recs.items():
        lst = []
        for f in rec["files"]:
            p = Path(files_dir) / f["file_name"]
            prob = None if p.is_file() and p.stat().st_size > 0 else "missing or empty"
            if not prob and p.read_bytes()[:5] != b"%PDF-": prob = "not a PDF"
            lst.append((p.name, p, prob))
        keys = [k for k, _, _ in lst]
        if len(keys) != len(set(keys)):
            lst = [(k, p, prob or "two files in this record have the same name") for k, p, prob in lst]
        out[rid] = lst
    return out


# ---------------------------------------------------------------- API client
class PrismError(Exception):
    pass

class Prism:
    def __init__(self, base_url, token, verify=True):
        import requests
        self.base = base_url.rstrip("/")
        self.s = requests.Session()
        self.s.headers.update({"Authorization": f"Bearer {token}", "Accept": "application/json"})
        self.verify = verify
        if verify is False:
            import urllib3; urllib3.disable_warnings()

    def call(self, method, path, json_body=None, data=None, content_type=None, ok=(200, 201, 202, 204), tries=3):
        url = path if path.startswith("http") else self.base + path
        headers = {"Content-Type": content_type} if content_type else ({"Content-Type": "application/json"} if json_body is not None else {})
        for attempt in range(tries):
            try:
                resp = self.s.request(method, url, json=json_body, data=data, headers=headers, verify=self.verify, timeout=120)
            except Exception as e:
                if attempt == tries - 1:
                    raise PrismError(f"{method} {path}: {e.__class__.__name__}: {e}")
                time.sleep(2 * (attempt + 1)); continue
            if resp.status_code in (429, 502, 503, 504) and attempt < tries - 1:
                time.sleep(3 * (attempt + 1)); continue
            body = None
            if resp.content:
                try: body = resp.json()
                except ValueError: body = resp.text[:500]
            if resp.status_code not in ok:
                msg = body.get("message") if isinstance(body, dict) else body
                errs = body.get("errors") if isinstance(body, dict) else None
                err = PrismError(f"{method} {path} -> HTTP {resp.status_code}: {msg}" + (f" {errs}" if errs else ""))
                err.status, err.body = resp.status_code, body
                raise err
            return resp.status_code, body

    def get(self, path, ok=(200,)):
        try:
            return self.call("GET", path, ok=ok)[1]
        except PrismError as e:
            if "HTTP 404" in str(e): return None
            raise

    def whoami(self):
        # an authenticated read that works for any user: their own records/drafts
        return self.call("GET", "/api/user/records?size=1")[1]


# ---------------------------------------------------------------- vocabularies (checked against Prism, never guessed)
# Prism's ids (staging, checked 2026-10-23). Bare ids such as "article" are category headings and are rejected.
RESOURCE_TYPE_GUESSES = {"report": ["text-report"],
                         "journal article": ["article-journal_article"],
                         "presentation": ["learning-presentation"]}

def resolve_vocab(api, recs, overrides=None):
    overrides = overrides or {}
    vocab = {"resource_types": {}, "licenses": {}, "roles": {}, "languages": {}, "notes": []}
    for t in sorted({_s(r["row"].get("Resource type")).lower() for r in recs.values()}):
        if overrides.get(t):
            vocab["resource_types"][t] = overrides[t]; continue
        found = None
        hits = (api.get(f"/api/vocabularies/resourcetypes?q={quote(t)}&size=50") or {}).get("hits", {}).get("hits", [])
        for h in hits:      # exact title match first (e.g. "Journal Article" -> article-journal_article)
            if _s((h.get("title") or {}).get("en")).lower() == t:
                found = h["id"]; break
        if not found:
            for cand in RESOURCE_TYPE_GUESSES.get(t, []):
                if api.get(f"/api/vocabularies/resourcetypes/{cand}"):
                    found = cand; break
        vocab["resource_types"][t] = found
        if not found: vocab["notes"].append(f"Resource type '{t}': no match. Set it in RESOURCE_TYPE_IDS.")
    for lic in sorted({_s(r["row"].get("Licenses: Title")) for r in recs.values()} - {""}):
        v = api.get(f"/api/vocabularies/licenses/{quote(lic)}")
        vocab["licenses"][lic] = v["id"] if v else None
        if not v: vocab["notes"].append(f"License '{lic}' is not in Prism's license list; it will be sent as a custom license.")
    roles = set()
    for r in recs.values():
        roles.update(_split(r["row"].get("Contributors: Role")))
    all_roles = (api.get("/api/vocabularies/contributorsroles?size=200") or {}).get("hits", {}).get("hits", [])
    vocab["all_roles"] = {h["id"]: _s((h.get("title") or {}).get("en")) for h in all_roles}
    for role in sorted(roles):
        v = api.get(f"/api/vocabularies/contributorsroles/{quote(role)}")
        # accept the id or the title, ignoring case, spaces, hyphens and the word "role"
        # ("Other", "Other role" and "role-other" all match the same entry)
        nrm = lambda x: re.sub(r"[^a-z]", "", str(x).lower()).replace("role", "")
        found = v["id"] if v else next((i for i, t in vocab["all_roles"].items()
                                        if nrm(role) in (nrm(i), nrm(t))), None)
        vocab["roles"][role] = found
        if not found:
            vocab["notes"].append(f"Contributor role '{role}' not found in Prism. Prism's roles: "
                                  + ", ".join(f"{t} ({i})" for i, t in vocab["all_roles"].items()))
    for lang in sorted({x for r in recs.values() for x in _split(r["row"].get("Languages"))}):
        v = api.get(f"/api/vocabularies/languages/{quote(lang)}")
        vocab["languages"][lang] = bool(v)
        if not v: vocab["notes"].append(f"Language '{lang}' not found in Prism.")
    return vocab

def offline_vocab(recs, overrides=None):
    """For a dry run without the network: assume the usual InvenioRDM ids (checked for real before sending)."""
    overrides = overrides or {}
    v = {"resource_types": {t: overrides.get(t) or RESOURCE_TYPE_GUESSES.get(t, [None])[0]
                            for t in {_s(r["row"].get("Resource type")).lower() for r in recs.values()}},
         "licenses": {}, "roles": {}, "languages": {}, "notes": ["Offline: vocabulary ids not checked against Prism yet."]}
    for r in recs.values():
        l = _s(r["row"].get("Licenses: Title"))
        if l: v["licenses"][l] = l
        for role in _split(r["row"].get("Contributors: Role")): v["roles"][role] = role
    return v


# ---------------------------------------------------------------- registry (atomic JSON)
class Registry:
    def __init__(self, path):
        self.path = Path(path); self.data = {}
        if self.path.exists():
            self.data = json.loads(self.path.read_text(encoding="utf-8"))

    def get(self, rid):
        return self.data.setdefault(rid, {})

    def save(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp = tempfile.mkstemp(dir=self.path.parent, suffix=".tmp")
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(self.data, f, indent=1, sort_keys=True)
        # Windows: antivirus, OneDrive or an open editor can briefly lock the file; wait and retry
        for attempt in range(20):
            try:
                os.replace(tmp, self.path); return
            except PermissionError:
                time.sleep(0.25 * (attempt + 1))
        try:
            self.path.write_text(json.dumps(self.data, indent=1, sort_keys=True), encoding="utf-8")
        finally:
            try: os.remove(tmp)
            except OSError: pass


class Log:
    """Appends one line per action. The file is opened only for each write, so it is never left locked."""
    def __init__(self, path):
        self.path = Path(path); self.path.parent.mkdir(parents=True, exist_ok=True)
        if not self.path.exists():
            self._write(["time", "record_id", "step", "status", "detail"])

    def _write(self, row):
        for attempt in range(10):
            try:
                with open(self.path, "a", newline="", encoding="utf-8") as f:
                    csv.writer(f).writerow(row)
                return
            except PermissionError:
                time.sleep(0.3 * (attempt + 1))
        print(f"  (could not write to {self.path.name}; is it open in Excel?)")

    def __call__(self, rid, step, status, detail=""):
        self._write([datetime.now().isoformat(timespec="seconds"), rid, step, status, detail])
        print(f"  {rid:<10} {step:<14} {status:<8} {detail}")


# ---------------------------------------------------------------- community
def ensure_community(api, slug, title, description="", send=False, log=print):
    c = api.get(f"/api/communities/{slug}")
    if c:
        log(f"Community already exists: {c['metadata']['title']}  (id {c['id']}, slug {slug})")
        return c
    if not send:
        log(f"Community '{slug}' does not exist yet. Set CREATE_COMMUNITY = True to create it.")
        return None
    body = {"slug": slug, "access": {"visibility": "public", "member_policy": "closed", "record_policy": "closed"},
            "metadata": {"title": title, "description": description, "type": {"id": "organization"}}}
    try:
        _, c = api.call("POST", "/api/communities", json_body=body)
    except PrismError as e:
        if "type" in str(e):          # instance without the 'organization' community type
            body["metadata"].pop("type"); _, c = api.call("POST", "/api/communities", json_body=body)
        else:
            raise
    log(f"Created community: {title}  (id {c['id']}, slug {slug}). You are its owner.")
    return c

def find_user(api, email, netid):
    """Best effort: Prism only finds users who have logged in at least once (and may hide emails)."""
    for q in (email, f'email:"{email}"', netid, f"username:{netid}"):
        if not q: continue
        res = api.get(f"/api/users?q={quote(q)}&size=5") or {}
        hits = res.get("hits", {}).get("hits", [])
        for h in hits:
            em = _s(h.get("email")).lower(); un = _s(h.get("username")).lower()
            if em == _s(email).lower() or (netid and un == _s(netid).lower()) or len(hits) == 1:
                return h
    return None

def invite_members(api, community_id, members, send=False, log=print):
    for m in members:
        who = f"{m['name']} ({m.get('netid') or m.get('email')}) as {m['role']}"
        u = find_user(api, m.get("email"), m.get("netid"))
        if not u:
            log(f"NOT FOUND: {who}. They need to log in to this Prism once; then re-run this cell. "
                "(Or invite them from the community's Members page.)")
            continue
        if not send:
            log(f"Found {who} (user id {u['id']}). Set ADD_MEMBERS = True to send the invitation."); continue
        body = {"members": [{"type": "user", "id": str(u["id"])}], "role": m["role"].lower(),
                "message": "You have been invited to the IVDRS & SUDORS community in Prism."}
        try:
            api.call("POST", f"/api/communities/{community_id}/invitations", json_body=body)
            log(f"Invited {who}. They accept the invitation from their Prism requests/notifications.")
        except PrismError as e:
            log(f"Could not invite {who}: {e}")


# ---------------------------------------------------------------- drafts and files
def _md5(path):
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""): h.update(chunk)
    return h.hexdigest()

def save_draft(api, rid, payload, reg, log, reserve_doi=True):
    st = reg.get(rid)
    if st.get("published"):
        log(rid, "draft", "skip", f"already published as {st['id']} (edit/new version is a separate step)"); return st
    current = api.get(f"/api/records/{st['id']}/draft") if st.get("id") else None
    if current:
        if not (payload.get("pids") or {}).get("doi") and (current.get("pids") or {}).get("doi"):
            payload = dict(payload, pids=current["pids"])     # keep the DOI Prism already reserved
        _, body = api.call("PUT", f"/api/records/{st['id']}/draft", json_body=payload)
        log(rid, "draft", "updated", st["id"])
    else:
        _, body = api.call("POST", "/api/records", json_body=payload)
        st.update({"id": body["id"], "created": datetime.now().isoformat(timespec="seconds")}); reg.save()
        log(rid, "draft", "created", body["id"])
    errs = body.get("errors") or []
    st["draft_errors"] = [f"{e.get('field')}: {e.get('messages')}" for e in errs]
    if errs: log(rid, "draft", "WARN", "; ".join(st["draft_errors"])[:400])
    if reserve_doi and "doi" not in (payload.get("pids") or {}) and not (body.get("pids") or {}).get("doi"):
        try:
            api.call("POST", f"/api/records/{st['id']}/draft/pids/doi")
            log(rid, "doi", "reserved", "Prism DOI reserved")
        except PrismError as e:
            log(rid, "doi", "WARN", f"could not reserve a DOI: {e}")
    reg.save()
    return st

def upload_files(api, rid, files, reg, log):
    st = reg.get(rid); did = st["id"]
    existing = {e["key"]: e for e in (api.get(f"/api/records/{did}/draft/files") or {}).get("entries", [])}
    done = st.setdefault("files", {})
    for key, path, prob in files:
        if prob:
            log(rid, "file", "SKIP", f"{key}: {prob}"); continue
        md5 = _md5(path)
        e = existing.get(key)
        if e and e.get("status") == "completed" and e.get("checksum") == f"md5:{md5}":
            done[key] = md5; log(rid, "file", "skip", f"{key} already uploaded"); continue
        if e:
            api.call("DELETE", f"/api/records/{did}/draft/files/{quote(key)}")
        for attempt in range(3):
            api.call("POST", f"/api/records/{did}/draft/files", json_body=[{"key": key}])
            try:
                with open(path, "rb") as fh:
                    api.call("PUT", f"/api/records/{did}/draft/files/{quote(key)}/content", data=fh,
                             content_type="application/octet-stream")
                break
            except PrismError as e:
                if "HTTP 404" not in str(e) or attempt == 2: raise
                log(rid, "file", "retry", f"{key}: Prism lost the file entry; starting it again")
                time.sleep(3)
        _, c = api.call("POST", f"/api/records/{did}/draft/files/{quote(key)}/commit")
        if c and c.get("checksum") and c["checksum"] != f"md5:{md5}":
            raise PrismError(f"{rid} {key}: checksum mismatch after upload")
        done[key] = md5; reg.save()
        log(rid, "file", "uploaded", f"{key} ({path.stat().st_size:,} bytes)")
    expected = {k for k, _, p in files if not p}
    st["files_complete"] = expected <= set(done) and bool(expected or not files); reg.save()

def publish(api, rid, reg, log, community_id=None, message="IVDRS & SUDORS deposit"):
    """Publish one draft. With a community: submit it to the community and accept the request
    (you are the community owner), which publishes the record into the community."""
    st = reg.get(rid)
    if st.get("published"):
        log(rid, "publish", "skip", f"already published: {st.get('doi') or st['id']}"); return
    if not st.get("id"):
        log(rid, "publish", "SKIP", "no draft yet"); return
    if not st.get("files_complete"):
        log(rid, "publish", "SKIP", "files are not all uploaded"); return
    did = st["id"]
    if community_id:
        api.call("PUT", f"/api/records/{did}/draft/review",
                 json_body={"receiver": {"community": community_id}, "type": "community-submission"})
        _, req = api.call("POST", f"/api/records/{did}/draft/actions/submit-review",
                          json_body={"payload": {"content": f"<p>{html.escape(message)}</p>", "format": "html"}})
        st["request_id"] = req.get("id"); reg.save()
        if req.get("status") != "accepted":
            api.call("POST", f"/api/requests/{req['id']}/actions/accept",
                     json_body={"payload": {"content": "<p>Accepted by the community owner.</p>", "format": "html"}})
    else:
        api.call("POST", f"/api/records/{did}/draft/actions/publish")
    rec = None
    for _ in range(10):
        rec = api.get(f"/api/records/{did}")
        if rec: break
        time.sleep(2)
    st["published"] = bool(rec)
    st["doi"] = ((rec or {}).get("pids") or {}).get("doi", {}).get("identifier")
    st["url"] = ((rec or {}).get("links") or {}).get("self_html")
    reg.save()
    log(rid, "publish", "done" if rec else "WARN", f"{st.get('doi') or ''} {st.get('url') or ''}".strip())


# ---------------------------------------------------------------- diagnose a rejected payload
ADD_PARTS = ["pids", "metadata.creators (full: ORCIDs + affiliations)", "metadata.description", "metadata.rights",
             "metadata.subjects", "metadata.languages", "metadata.publisher", "metadata.identifiers",
             "metadata.contributors", "files", "custom_fields"]

def _minimal(payload):
    md = payload["metadata"]
    creators = [{"person_or_org": {k: v for k, v in c["person_or_org"].items() if k != "identifiers"}}
                for c in md.get("creators", [])]
    return {"access": payload["access"], "files": {"enabled": False}, "pids": {},
            "metadata": {"resource_type": md.get("resource_type"), "title": md.get("title"),
                         "publication_date": md.get("publication_date"), "creators": creators}}

def _with(base, payload, part):
    p = json.loads(json.dumps(base)); key = part.split(" ")[0]
    if key.startswith("metadata."):
        k = key.split(".", 1)[1]
        if k not in payload["metadata"]: return None
        p["metadata"][k] = payload["metadata"][k]
    else:
        if key not in payload: return None
        p[key] = payload[key]
    return p

def diagnose(api, rid, payload, log=print):
    """Start from a bare-bones version of the record (resource type, title, date, author names) and add one
    part at a time, so it finds the problem even when more than one part is rejected. Test drafts that Prism
    accepts are deleted straight away; the registry is not touched."""
    def attempt(p):
        try:
            _, body = api.call("POST", "/api/records", json_body=p)
        except PrismError as e:
            return False, getattr(e, "body", str(e))
        try:
            api.call("DELETE", f"/api/records/{body['id']}/draft")
        except PrismError as e:
            log(f"   (could not delete test draft {body['id']}: {e}; delete it from My dashboard)")
        return True, body.get("errors")
    base = _minimal(payload)
    ok, info = attempt(base)
    log(f"{rid}: bare-bones record (type, title, date, author names) -> {'accepted' if ok else 'REJECTED'}")
    if not ok:
        log("   response: " + json.dumps(info, ensure_ascii=False)[:800])
        for label, change in [("as resource type publication-report", ("resource_type", {"id": "publication-report"})),
                              ("with year-only date", ("publication_date", str(base["metadata"]["publication_date"])[:4])),
                              ("with only the first author", ("creators", base["metadata"]["creators"][:1]))]:
            t = json.loads(json.dumps(base)); t["metadata"][change[0]] = change[1]
            ok2, _ = attempt(t)
            log(f"   bare-bones {label:<38} -> {'accepted' if ok2 else 'rejected'}")
        log("   payload sent: " + json.dumps(base, ensure_ascii=False)[:1500])
        return ["base"]
    culprits = []
    for part in ADD_PARTS:
        trial = _with(base, payload, part)
        if trial is None: continue
        ok, info = attempt(trial)
        log(f"   + {part:<46} -> {'accepted' if ok else 'REJECTED'}")
        if not ok:
            culprits.append(part)
            key = part.split(" ")[0]
            val = payload.get(key) if not key.startswith("metadata.") else payload["metadata"].get(key.split(".", 1)[1])
            log("       value sent: " + json.dumps(val, ensure_ascii=False)[:700])
    log(f"=> Prism rejects: {', '.join(culprits) if culprits else 'nothing on its own; send Claude this output'}")
    return culprits


def diagnose_from_good(api, good_rid, good_payload, bad_rid, bad_payload, log=print):
    """Start from a payload Prism accepted (a report that worked) and change one thing at a time toward the
    rejected one. Test drafts Prism accepts are deleted straight away; the registry is not touched."""
    def attempt(p):
        try:
            _, body = api.call("POST", "/api/records", json_body=p)
        except PrismError as e:
            return False, getattr(e, "body", str(e))
        try: api.call("DELETE", f"/api/records/{body['id']}/draft")
        except PrismError as e: log(f"   (could not delete test draft {body['id']}: {e})")
        return True, None
    g = json.loads(json.dumps(good_payload)); b = bad_payload["metadata"]
    ok, info = attempt(g)
    log(f"{good_rid} as sent (worked before) -> {'accepted' if ok else 'REJECTED ' + json.dumps(info)}")
    first = {"person_or_org": {k: v for k, v in b["creators"][0]["person_or_org"].items() if k != "identifiers"}}
    steps = [
        ("resource type of the article " + json.dumps(b.get("resource_type")), lambda p: p["metadata"].__setitem__("resource_type", b["resource_type"])),
        ("one personal author (no ORCID, no affiliation)", lambda p: p["metadata"].__setitem__("creators", [first])),
        ("personal author with given/family as 'Test', 'Person'", lambda p: p["metadata"].__setitem__("creators", [{"person_or_org": {"type": "personal", "given_name": "Test", "family_name": "Person"}}])),
        ("files disabled", lambda p: p.__setitem__("files", {"enabled": False})),
        ("the article's title", lambda p: p["metadata"].__setitem__("title", b["title"])),
        ("the article's date", lambda p: p["metadata"].__setitem__("publication_date", b["publication_date"])),
        ("the article's pids " + json.dumps(bad_payload.get("pids")), lambda p: p.__setitem__("pids", bad_payload.get("pids", {}))),
    ]
    for label, change in steps:
        t = json.loads(json.dumps(g)); change(t)
        ok, info = attempt(t)
        log(f"   change only: {label:<70} -> {'accepted' if ok else 'REJECTED'}")
    rts = api.get("/api/vocabularies/resourcetypes?size=200") or {}
    hits = rts.get("hits", {}).get("hits", [])
    log(f"\nResource types Prism lists ({len(hits)}):")
    for h in hits:
        log(f"   {h.get('id'):<28} {(h.get('title') or {}).get('en', '')}")


PRODUCTION_PHRASE = "DEPOSIT TO PRODUCTION"

def require_confirmation(env, typed):
    """Every step that changes Prism production stops unless the phrase was typed in the settings cell."""
    if env == "production" and (typed or "").strip() != PRODUCTION_PHRASE:
        raise SystemExit(f'Stopped: this changes Prism PRODUCTION. Type CONFIRM_PRODUCTION = "{PRODUCTION_PHRASE}" '
                         "in cell 2, run cell 2 again, then re-run this cell.")


# ---------------------------------------------------------------- community records report
CHILD_TABS = ["IVDRS Reports and Articles", "SUDORS Reports and Articles", "Peer-reviewed Publications"]

def community_records(api, community_id, log=print):
    """Every published record in the community (all pages)."""
    out, page = [], 1
    while True:
        res = api.get(f"/api/communities/{community_id}/records?size=100&page={page}&sort=newest") or {}
        hits = res.get("hits", {}).get("hits", [])
        out += hits
        if not hits or len(out) >= res.get("hits", {}).get("total", 0):
            break
        page += 1
    log(f"{len(out)} published records in the community")
    return out

def _doi_link(d):
    return f"https://doi.org/{d}" if d else None

def report_rows(records, registry=None):
    by_id = {v.get("id"): rid for rid, v in (registry.data.items() if registry else []) if v.get("id")}
    rows = []
    for r in records:
        md = r.get("metadata", {}); parent = r.get("parent", {}) or {}
        parent_doi = ((parent.get("pids") or {}).get("doi") or {}).get("identifier")
        version_doi = ((r.get("pids") or {}).get("doi") or {}).get("identifier")
        provider = ((r.get("pids") or {}).get("doi") or {}).get("provider")
        creators = "; ".join(c.get("person_or_org", {}).get("name") or
                             ", ".join(x for x in [c.get("person_or_org", {}).get("family_name"),
                                                   c.get("person_or_org", {}).get("given_name")] if x)
                             for c in md.get("creators", []))
        subjects = [s.get("subject") or (s.get("id") or "") for s in md.get("subjects", [])]
        cite = parent_doi or version_doi
        if parent_doi:
            note = "Cite the parent DOI: it stays the same when new versions are added."
        elif provider == "external":
            note = "Publisher's DOI (Prism did not mint one). It does not change between Prism versions."
        else:
            note = "No parent DOI returned; cite the version DOI for now."
        rows.append({
            "record_id": by_id.get(r.get("id")) or by_id.get(parent.get("id")),
            "Title": md.get("title"),
            "Resource type": (md.get("resource_type") or {}).get("title", {}).get("en") or (md.get("resource_type") or {}).get("id"),
            "Publication date": md.get("publication_date"),
            "Creators": creators,
            "DOI to cite": cite, "DOI to cite (link)": _doi_link(cite),
            "Parent DOI": parent_doi, "Version DOI": version_doi,
            "Version": (r.get("versions") or {}).get("index"),
            "Prism record": (r.get("links") or {}).get("self_html"),
            "Files": len(((r.get("files") or {}).get("entries") or {})) or (r.get("files") or {}).get("count"),
            "License": "; ".join(x.get("id") or (x.get("title") or {}).get("en", "") for x in md.get("rights", [])),
            "Note": note,
            "_subjects": subjects,
            "_external": provider == "external",
        })
    rows.sort(key=lambda x: (x["record_id"] or "~", x["Title"] or ""))
    return rows

def save_report(rows, path, community_title):
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    cols = [c for c in rows[0] if not c.startswith("_")] if rows else []
    head = Font(name="Arial", bold=True, color="FFFFFF", size=10); fill = PatternFill("solid", fgColor="4E2A84")
    body = Font(name="Arial", size=10); link = Font(name="Arial", size=10, color="0563C1", underline="single")
    widths = {"record_id": 11, "Title": 60, "Resource type": 15, "Publication date": 11, "Creators": 40,
              "DOI to cite": 28, "DOI to cite (link)": 38, "Parent DOI": 28, "Version DOI": 28, "Version": 8,
              "Prism record": 40, "Files": 6, "License": 16, "Note": 50}
    tabs = [("All records", rows)] + [(t, [r for r in rows if t in r["_subjects"]]) for t in CHILD_TABS]
    for name, rs in tabs:
        ws = wb.create_sheet(name[:31])
        for c, h in enumerate(cols, 1):
            x = ws.cell(1, c, h); x.font = head; x.fill = fill; x.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = widths.get(h, 14)
        for i, r in enumerate(rs, 2):
            for c, h in enumerate(cols, 1):
                x = ws.cell(i, c, r[h]); x.font = body; x.alignment = Alignment(wrap_text=True, vertical="top")
                if h in ("DOI to cite (link)", "Prism record") and r[h]:
                    x.hyperlink = r[h]; x.font = link
        ws.freeze_panes = "C2"
        if rs: ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(cols))}{len(rs) + 1}"
    ws = wb.create_sheet("About", 0)
    notes = [community_title, "",
             "One row per record. 'All records' is everything in the parent community; the other tabs are the proposed "
             "child communities (a record can be on more than one: IVDRS- and SUDORS-related articles also appear under "
             "Peer-reviewed Publications).",
             "Please cite the 'DOI to cite'. For reports this is the parent DOI, which always points to the latest "
             "version and does not change when a corrected file is added as a new version. (Prism's 'Cite as' box "
             "shows the version DOI, which belongs to one version only.)",
             "Journal articles keep the publisher's DOI."]
    for i, t in enumerate(notes, 1):
        x = ws.cell(i, 1, t); x.font = Font(name="Arial", size=12 if i == 1 else 10, bold=i == 1)
        x.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 110
    for i, (name, rs) in enumerate(tabs, len(notes) + 2):
        ws.cell(i, 1, f"{name}: {len(rs)} records").font = Font(name="Arial", size=10)
    wb.save(path)
    return {n: len(rs) for n, rs in tabs}
